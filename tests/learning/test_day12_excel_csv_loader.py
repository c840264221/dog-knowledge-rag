from datetime import date, datetime

import csv
from pathlib import Path

from tests.learning.test_day08_unified_document import (
    DocumentBlock,
    SourceLocation,
    UnifiedDocument,
)
from collections.abc import Iterable, Sequence

from openpyxl import load_workbook
from tests.learning.test_day09_parser_registry import (
    DocumentParser,
    DocumentSourceType,
    parse_document
)

from tests.learning.test_day10_pdf_loader import TextPdfParser

def normalize_table_cell_value(
    value: object,
) -> str:
    """
    将 CSV 或 Excel 单元格的值转换成适合检索的文字。

    功能：
        统一处理表格中可能出现的空值、字符串、日期、
        日期时间、数字和布尔值。

        空值转换成空字符串；
        普通字符串去掉前后空格；
        日期转换成“年-月-日”；
        日期时间转换成稳定的日期时间文字；
        其他类型使用 str() 转换后再去掉前后空格。

    参数：
        value:
            从 CSV 或 Excel 单元格中读取到的原始值。
            可能是 None、字符串、数字、布尔值、
            date 或 datetime。

    返回值：
        str:
            清理并转换后的单元格文字。
            空单元格返回空字符串。
    """
    # normalized_value = None
    # 请在这里补充代码
    if value is None:
        return ""

    if isinstance(value, datetime):
        return value.isoformat(
            sep=" ",
            timespec="seconds",
        )

    if isinstance(value, date):
        return value.isoformat()

    return str(value).strip()

def normalize_table_headers(
    headers: list[object],
) -> list[str]:
    """
    清理并校验 CSV 或 Excel 的表头。

    功能：
        使用 normalize_table_cell_value 清理每个表头，
        检查表头列表不能为空，
        检查每个表头都包含有效文字，
        并拒绝清理后重复的表头名称。

    参数：
        headers:
            从 CSV 第一行或 Excel 工作表表头行中
            读取到的原始单元格值。

    返回值：
        list[str]:
            保持原有列顺序的有效表头文字。

    异常：
        ValueError:
            当表头列表为空、某个表头为空，
            或者存在重复表头时抛出。
    """
    # 请在这里补充代码
    if len(headers) == 0:
        raise ValueError("不能没有表头数据")

    seen_headers: set[str] = set()
    normalized_headers: list[str] = []

    for header in headers:
        normalized_header = normalize_table_cell_value(header)

        if not normalized_header:
            raise ValueError("清理后的表头字符为空")

        if normalized_header not in seen_headers:
            normalized_headers.append(normalized_header)
            seen_headers.add(normalized_header)
        else:
            raise ValueError("表头重复")

    return normalized_headers

def format_table_row(
    *,
    headers: list[str],
    row_values: list[object],
) -> str | None:
    """
    将一行 CSV 或 Excel 数据转换成适合检索的文字。

    功能：
        按照列顺序把表头和当前行的单元格值一一对应，
        使用 normalize_table_cell_value 清理每个值，
        跳过内容为空的单元格，
        最后把有效的“表头：值”组合成一段文字。

        如果这一行的所有单元格都是空的，则返回 None。

        本函数接收的 headers 应该已经通过
        normalize_table_headers 完成清理和校验。

    参数：
        headers:
            已经清理完成的表头列表。

        row_values:
            当前数据行中各个单元格的原始值。

    返回值：
        str | None:
            有有效数据时返回整理后的行文字；
            整行为空时返回 None。

            例如：
            "犬种：柴犬 | 原产地：日本 | 体型：中型"

    异常：
        ValueError:
            当表头数量和当前行的单元格数量不一致时抛出。
            这样可以避免表头和值发生错误对应。
    """
    # 请在这里补充代码
    if len(headers) != len(row_values):
        raise ValueError("表头数量和行内容数量不一致")

    row_parts: list[str] = []
    for header, row_value in zip(headers, row_values):
        normalized_value = normalize_table_cell_value(row_value)
        if not normalized_value:
            continue

        formatted_row = f"{header}：{normalized_value}"
        row_parts.append(formatted_row)

    if not row_parts:
        return None

    return " | ".join(row_parts)

def build_table_row_block(
    *,
    source: str,
    sheet_name: str | None,
    row_number: int,
    block_index: int,
    headers: list[str],
    row_values: list[object],
) -> DocumentBlock | None:
    """
    将一行 CSV 或 Excel 数据转换成表格内容块。

    功能：
        校验源文件、工作表名称、原始行号和内容块顺序，
        调用 format_table_row 生成适合检索的行文字，
        并创建包含原文件位置的 DocumentBlock。

        CSV 没有工作表概念，因此 sheet_name 可以为 None。
        Excel 数据需要保存真实工作表名称。

        如果当前数据行没有任何有效内容，
        则返回 None，不创建空的 DocumentBlock。

    参数：
        source:
            CSV 或 Excel 原文件路径，不能为空。

        sheet_name:
            Excel 工作表名称。
            CSV 数据可以传入 None。
            如果传入字符串，则不能为空白字符串。

        row_number:
            当前数据在原始表格中的真实行号，从 1 开始。
            通常第 1 行是表头，因此第一行数据可能是第 2 行。

        block_index:
            当前内容块在统一文档中的顺序，从 0 开始。

        headers:
            已经完成清理和校验的表头列表。

        row_values:
            当前数据行中的原始单元格值。

    返回值：
        DocumentBlock | None:
            当前行存在有效内容时，返回 table 类型的内容块；
            当前行所有单元格都为空时，返回 None。

    异常：
        ValueError:
            当 source 为空、sheet_name 是空白字符串、
            row_number 小于 1、block_index 小于 0，
            或表头和值数量不一致时抛出。
    """
    # 请在这里补充代码
    normalized_source = source.strip()
    if not normalized_source:
        raise ValueError("表格来源不能为空")

    normalized_sheet_name: str | None = None
    if sheet_name is not None:
        normalized_sheet_name = sheet_name.strip()
        if not normalized_sheet_name:
            raise ValueError("表格sheet name不能为空")

    if row_number < 1:
        raise ValueError("行号不能小于1")

    if block_index < 0:
        raise ValueError("内容块的顺序号不能小于0")

    formatted_row = format_table_row(headers=headers, row_values=row_values)

    if formatted_row is None:
        return None

    location = SourceLocation(
        source=normalized_source,
        sheet_name=normalized_sheet_name,
        row_start=row_number,
        row_end=row_number,
    )
    block_id = f"table-row-{block_index}"

    return DocumentBlock(
        block_id=block_id,
        block_type="table",
        content=formatted_row,
        block_index=block_index,
        location=location,
        metadata={
            "column_count": len(headers),
        },
    )

def load_csv_document(
    file_path: str,
    encoding: str = "utf-8-sig",
) -> UnifiedDocument:
    """
    读取 CSV 文件并转换成统一文档。

    功能：
        检查 CSV 文件路径和字符编码，
        使用 Python 自带的 csv 模块读取文件，
        把第一行作为表头进行清理和校验，
        再把后续每一行有效数据转换成 DocumentBlock，
        最后组成 UnifiedDocument。

        CSV 没有工作表，因此生成的 SourceLocation
        中 sheet_name 保存为 None。

        空数据行不会生成 DocumentBlock，
        但会记录在文档 metadata 中。

    参数：
        file_path:
            需要读取的 CSV 文件路径。

        encoding:
            CSV 文件使用的字符编码。
            默认使用 utf-8-sig，
            它既能读取普通 UTF-8，也能处理带 BOM 的 UTF-8 文件。

    返回值：
        UnifiedDocument:
            包含 CSV 基本信息、表头、有效数据行 Block
            和空行诊断信息的统一文档。

    异常：
        ValueError:
            当路径为空、目标不是 CSV 文件、路径不是普通文件、
            编码名称为空、CSV 没有表头，
            或者数据行列数与表头数量不一致时抛出。

        FileNotFoundError:
            当 CSV 文件不存在时抛出。

        UnicodeDecodeError:
            当指定编码无法正确读取文件时抛出。

        csv.Error:
            当 CSV 内容格式损坏时，由 csv 模块向外抛出。
    """
    # 请在这里补充代码
    normalized_path = file_path.strip()
    if not normalized_path:
        raise ValueError("表格文件路径不能为空")

    csv_path = Path(normalized_path)
    if not csv_path.exists():
        raise FileNotFoundError(
            f"CSV 文件不存在：{csv_path}"
        )

    if not csv_path.is_file():
        raise ValueError(f"{normalized_path}不是文件类型")

    if csv_path.suffix.lower() != ".csv":
        raise ValueError(f"{normalized_path}文件不是csv格式")

    normalized_encoding = encoding.strip().lower()
    if not normalized_encoding:
        raise ValueError("编码不能为空")

    with csv_path.open(
            mode="r",
            encoding=normalized_encoding,
            newline="",
    ) as csv_file:
        csv_reader = csv.reader(csv_file)

        raw_headers = next(csv_reader, None)

        if raw_headers is None:
            raise ValueError("CSV 文件没有表头")

        headers = normalize_table_headers(raw_headers)

        row_blocks: list[DocumentBlock] = []
        empty_row_numbers: list[int] = []

        for row_number, row_values in enumerate(
                csv_reader,
                start=2,
        ):
            row_is_empty = (
                not row_values
                or all(
                    not normalize_table_cell_value(value)
                    for value in row_values
                )
            )

            if row_is_empty:
                empty_row_numbers.append(row_number)
                continue

            row_block = build_table_row_block(
                source=normalized_path,
                sheet_name=None,
                row_number=row_number,
                block_index=len(row_blocks),
                headers=headers,
                row_values=row_values,
            )
            if row_block is None:
                empty_row_numbers.append(row_number)
                continue

            row_blocks.append(row_block)

        return UnifiedDocument(
            doc_id=f"CSV-{csv_path.stem}",
            source=str(csv_path),
            source_type="csv",
            title=csv_path.stem,
            blocks=row_blocks,
            metadata={
                "parser": "csv",
                "encoding": normalized_encoding,
                "headers": headers,
                "column_count": len(headers),
                "data_row_count": len(row_blocks) + len(empty_row_numbers),
                "block_count": len(row_blocks),
                "empty_row_numbers": empty_row_numbers,
            },
        )

def run_csv_loader_demo(
    *,
    file_path: str,
    encoding: str = "utf-8-sig",
) -> None:
    """
    演示真实 CSV 文件转换成统一文档的过程。

    功能：
        调用 load_csv_document 读取真实 CSV，
        打印统一文档的基本信息和诊断数据，
        再按照 Block 顺序打印每一行有效数据的
        原始行号和可检索文字。

        本函数只用于学习和手动检查，
        不属于正式 RAG 业务代码。

    参数：
        file_path:
            需要读取的 CSV 文件路径。

        encoding:
            CSV 文件使用的字符编码，
            默认使用 utf-8-sig。

    返回值：
        None:
            不返回业务数据，只把读取结果打印到控制台。

    输出格式：
        文档编号
        文件类型
        表头
        数据总行数
        有效 Block 数量
        空行行号
        每个 Block 的编号、原始行号和正文
    """
    # 请在这里补充代码
    document = load_csv_document(
        file_path=file_path,
        encoding=encoding,
    )
    print(f"文档编号：{document.doc_id}")
    print(f"文件类型：{document.source_type}")
    print(f"表头：{document.metadata['headers']}")
    print(f"数据总行数：{document.metadata['data_row_count']}")
    print(f"有效 Block 数量：{document.metadata['block_count']}")
    print(f"空行行号：{document.metadata['empty_row_numbers']}")

    for block in document.blocks:
        location = block.location

        print(f"Block：{block.block_id}")
        print(f"原始行号：{location.row_start}")
        print(f"正文：{block.content}")
        print("-" * 50)

def build_excel_sheet_blocks(
    *,
    source: str,
    sheet_name: str,
    rows: Iterable[Sequence[object]],
    start_block_index: int,
) -> list[DocumentBlock]:
    """
    将一个 Excel 工作表的行数据转换成表格内容块。

    功能：
        把工作表第一行作为表头进行清理和校验，
        再从真实行号 2 开始遍历后续数据行，
        将每一行有效数据转换成 DocumentBlock。

        空数据行不会生成 Block。

        start_block_index 用于保证多个工作表生成的
        DocumentBlock 顺序不会重新从 0 开始。

    参数：
        source:
            Excel 原文件路径，不能为空。

        sheet_name:
            当前 Excel 工作表名称，不能为空。

        rows:
            当前工作表的所有行。
            第一行必须是表头，后续行为数据。
            通常来自 worksheet.iter_rows(values_only=True)。

        start_block_index:
            当前工作表第一个 Block 使用的顺序号，
            必须大于或等于 0。

    返回值：
        list[DocumentBlock]:
            当前工作表生成的有效表格内容块。
            如果工作表只有表头而没有数据，返回空列表。

    异常：
        ValueError:
            当 source 为空、sheet_name 为空、
            start_block_index 小于 0、工作表完全为空、
            表头不合法，或者数据行列数不正确时抛出。
    """
    # 请在这里补充代码
    normalized_source = source.strip()
    if not normalized_source:
        raise ValueError("文件来源不能为空")

    normalized_sheet_name = sheet_name.strip()
    if not normalized_sheet_name:
        raise ValueError("工作表名不能为空")

    if start_block_index < 0:
        raise ValueError("起始 Block 顺序号不能小于 0")

    row_iterator = iter(rows)
    raw_headers = next(row_iterator, None)

    if raw_headers is None:
        raise ValueError("工作表没有内容")


    headers = normalize_table_headers(list(raw_headers))

    sheet_blocks: list[DocumentBlock] = []

    for row_number, row_values in enumerate(
            row_iterator,
            start=2,
    ):
        block_index = start_block_index + len(sheet_blocks)
        row_block = build_table_row_block(
            source=normalized_source,
            sheet_name=normalized_sheet_name,
            row_number=row_number,
            block_index=block_index,
            headers=headers,
            row_values=list(row_values),
        )
        if row_block is None:
            continue

        sheet_blocks.append(row_block)

    return sheet_blocks

def load_excel_document(
    file_path: str,
) -> UnifiedDocument:
    """
    读取 Excel 工作簿并转换成统一文档。

    功能：
        检查 Excel 文件路径，
        使用 openpyxl 以只读方式打开 .xlsx 文件，
        遍历工作簿中的所有工作表，
        把每个工作表的数据转换成 DocumentBlock，
        最后组成一个 UnifiedDocument。

        不同工作表生成的 Block 使用连续的 block_index，
        不会在每个工作表中重新从 0 开始。

        当前基础版本要求每个工作表都有有效表头，
        暂时不处理完全空白的工作表。

    参数：
        file_path:
            需要读取的 .xlsx 文件路径。

    返回值：
        UnifiedDocument:
            包含全部工作表数据、Block 和工作表统计信息的
            统一 Excel 文档。

    异常：
        FileNotFoundError:
            当 Excel 文件不存在时抛出。

        ValueError:
            当路径为空、路径不是普通文件、文件不是 .xlsx、
            工作簿没有工作表、工作表为空、表头不合法，
            或数据行列数错误时抛出。

        openpyxl 相关异常:
            当 Excel 文件损坏、加密或不是合法工作簿时，
            保留 openpyxl 原始异常向外抛出。
    """
    # 请在这里补充代码

    normalized_path = file_path.strip()
    if not normalized_path:
        raise ValueError("表格文件路径不能为空")

    xlsx_path = Path(normalized_path)
    if not xlsx_path.exists():
        raise FileNotFoundError(
            f"Excel 文件不存在：{xlsx_path}"
        )

    if not xlsx_path.is_file():
        raise ValueError(f"{normalized_path}不是文件类型")

    if xlsx_path.suffix.lower() != ".xlsx":
        raise ValueError("当前 Excel Loader 只支持 .xlsx 文件")


    workbook = load_workbook(
        filename=normalized_path,
        read_only=True,
        data_only=True,
    )
    all_blocks: list[DocumentBlock] = []
    sheet_names: list[str] = []
    sheet_block_counts: dict[str, int] = {}
    try:
        if not workbook.worksheets:
            raise ValueError("Excel 工作簿没有工作表")

        for worksheet in workbook.worksheets:
            sheet_blocks = build_excel_sheet_blocks(
                source=normalized_path,
                sheet_name=worksheet.title,
                rows=worksheet.iter_rows(values_only=True),
                start_block_index=len(all_blocks),
            )
            sheet_names.append(worksheet.title)
            sheet_block_counts[worksheet.title] = len(sheet_blocks)
            all_blocks.extend(sheet_blocks)
    finally:
        workbook.close()

    document = UnifiedDocument(
        doc_id=f"EXCEL-{xlsx_path.stem}",
        source=normalized_path,
        source_type="excel",
        title=xlsx_path.stem,
        blocks=all_blocks,
        metadata={
            "parser": "openpyxl",
            "sheet_count": len(sheet_names),
            "sheet_names": sheet_names,
            "sheet_block_counts": sheet_block_counts,
            "block_count": len(all_blocks),
            "read_only": True,
            "data_only": True,
        },
    )
    return document

def run_excel_loader_demo(
    *,
    file_path: str,
) -> None:
    """
    演示真实 Excel 工作簿转换成统一文档的过程。

    功能：
        调用 load_excel_document 读取真实 .xlsx 文件，
        打印工作簿的基本信息、工作表列表和 Block 数量，
        再按照统一顺序打印每个 Block 的工作表名称、
        原始行号和可检索正文。

        本函数只用于学习和手动检查，
        不属于正式 RAG 业务代码。

    参数：
        file_path:
            需要读取的 .xlsx 文件路径。

    返回值：
        None:
            不返回业务数据，只把解析结果打印到控制台。

    输出格式：
        文档编号
        文件类型
        工作表数量
        工作表名称
        每个工作表的 Block 数量
        全部 Block 的编号、工作表、原始行号和正文
    """
    # 请在这里补充代码
    document = load_excel_document(file_path)
    print(f"文档编号：{document.doc_id}")
    print(f"文件类型：{document.source_type}")
    print(f"工作表数量：{document.metadata['sheet_count']}")
    print(f"工作表名称：{document.metadata['sheet_names']}")
    print(
        "工作表 Block 数量："
        f"{document.metadata['sheet_block_counts']}"
    )
    print(f"全部 Block 数量：{document.metadata['block_count']}")

    for block in document.blocks:
        location = block.location

        print(f"Block：{block.block_id}")
        print(f"工作表：{location.sheet_name}")
        print(f"原始行号：{location.row_start}")
        print(f"正文：{block.content}")
        print("-" * 50)

class CsvDocumentParser:
    """
    CSV 统一文档解析器。

    功能：
        实现 DocumentParser 约定的 parse 方法，
        把统一解析入口传入的 CSV 文件路径，
        交给 load_csv_document 完成实际读取。

    返回值：
        UnifiedDocument:
            source_type 为 csv 的统一文档。
    """

    def parse(
        self,
        file_path: str,
    ) -> UnifiedDocument:
        """
        解析 CSV 文件并返回统一文档。

        参数：
            file_path:
                需要解析的 CSV 文件路径。

        返回值：
            UnifiedDocument:
                CSV Loader 生成的统一文档。
        """
        # 请补充代码
        return load_csv_document(file_path)

class ExcelDocumentParser:
    """
    Excel 统一文档解析器。

    功能：
        实现 DocumentParser 约定的 parse 方法，
        把统一解析入口传入的 .xlsx 文件路径，
        交给 load_excel_document 完成实际读取。

        当前版本只支持 .xlsx，不支持旧版 .xls。

    返回值：
        UnifiedDocument:
            source_type 为 excel 的统一文档。
    """

    def parse(
        self,
        file_path: str,
    ) -> UnifiedDocument:
        """
        解析 Excel 文件并返回统一文档。

        参数：
            file_path:
                需要解析的 .xlsx 文件路径。

        返回值：
            UnifiedDocument:
                Excel Loader 生成的统一文档。
        """
        # 请补充代码
        return load_excel_document(file_path)

def build_multiformat_parser_registry(
) -> dict[DocumentSourceType, DocumentParser]:
    """
    创建当前学习阶段的多格式文档解析器登记表。

    功能：
        登记已经真正实现的 PDF、CSV 和 Excel 解析器，
        让统一解析入口可以根据文件扩展名选择对应 Parser。

        当前还没有实现 Markdown Parser，
        因此不提前登记 markdown。

    参数：
        无。

    返回值：
        dict[DocumentSourceType, DocumentParser]:
            多格式解析器登记表。

            键表示统一文档类型：
                pdf
                csv
                excel

            值表示对应的解析器对象：
                TextPdfParser
                CsvDocumentParser
                ExcelDocumentParser
    """
    # 请在这里补充代码
    return {
        "pdf": TextPdfParser(),
        "excel": ExcelDocumentParser(),
        "csv": CsvDocumentParser(),
    }

def run_multiformat_parser_demo(
    file_paths: Sequence[str],
) -> None:
    """
    使用统一解析入口依次解析多种格式的文档。

    功能：
        创建多格式解析器注册表，
        依次把传入的文件路径交给 parse_document。
        parse_document 会根据文件扩展名自动选择 PDF、CSV 或 Excel 解析器，
        最后打印每份统一文档的基本信息。

    参数：
        file_paths:
            等待解析的文件路径集合。
            可以同时包含 PDF、CSV 和 Excel 文件。

    返回值：
        None：
            本函数只打印解析结果，不返回其他数据。
    """
    # 具体逻辑由你完成
    if not file_paths:
        raise ValueError("文件路径列表不能为空")
    parser_registry = build_multiformat_parser_registry()
    for file_path in file_paths:
        parsed_document = parse_document(
            file_path=file_path,
            parser_registry=parser_registry,
        )

        print(f"文件路径：{file_path}")
        print(f"文档编号：{parsed_document.doc_id}")
        print(f"文件类型：{parsed_document.source_type}")
        print(f"Block 数量：{len(parsed_document.blocks)}")
        print("-" * 50)


if __name__ == "__main__":
    run_multiformat_parser_demo(
        file_paths=[
            "output/pdf/day10_text_sample.pdf",
            "output/csv/day12_dog_breeds.csv",
            "output/excel/day12_dog_tables.xlsx",
        ],
    )